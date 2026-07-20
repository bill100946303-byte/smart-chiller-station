from __future__ import annotations

from datetime import date
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


OUT = Path("docs/plc/冷站PLC标准化编程与AI系统直写实施说明.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(28, 35, 43)
MUTED = RGBColor(90, 98, 110)
RISK = RGBColor(155, 28, 28)
GOOD = RGBColor(24, 112, 72)
CAUTION = RGBColor(122, 90, 0)

FILL_BLUE = "E8EEF5"
FILL_GRAY = "F2F4F7"
FILL_CODE = "F7F9FB"
FILL_WARN = "FFF4E5"
FILL_GOOD = "EAF7F0"
FILL_RED = "FDECEC"

LIST_NUM_ID = 41
BULLET_NUM_ID = 42
LIST_COMMAND_NUM_ID = 43
LIST_MANUAL_NUM_ID = 44
LIST_B25_NUM_ID = 45

DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def set_font(
    run,
    name="Arial",
    east_asia="PingFang SC",
    size=None,
    color=None,
    bold=None,
    italic=None,
):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_para_spacing(paragraph, before=0, after=6, line=1.25):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(index, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def paragraph_border_bottom(paragraph, color="2E74B5", size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MUTED)


def add_numbering_definition(doc):
    numbering = doc.part.numbering_part.element

    def add_abstract(abstract_id, num_fmt, level_text, font_name=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)

        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), level_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        level.extend([start, fmt, text, suff])

        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.extend([tabs, ind, spacing])
        level.append(p_pr)

        if font_name:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font_name)
            fonts.set(qn("w:hAnsi"), font_name)
            r_pr.append(fonts)
            level.append(r_pr)

        abstract.append(level)
        numbering.append(abstract)

    add_abstract(91, "decimal", "%1.")
    add_abstract(92, "bullet", "•", "Symbol")
    add_abstract(93, "decimal", "%1.")
    add_abstract(94, "decimal", "%1.")
    add_abstract(95, "decimal", "%1.")

    for num_id, abstract_id in (
        (LIST_NUM_ID, 91),
        (BULLET_NUM_ID, 92),
        (LIST_COMMAND_NUM_ID, 93),
        (LIST_MANUAL_NUM_ID, 94),
        (LIST_B25_NUM_ID, 95),
    ):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_num_id = OxmlElement("w:abstractNumId")
        abstract_num_id.set(qn("w:val"), str(abstract_id))
        num.append(abstract_num_id)
        numbering.append(num)


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def add_p(doc, text="", bold=False, color=INK, size=11, after=6, keep=False):
    paragraph = doc.add_paragraph()
    set_para_spacing(paragraph, after=after)
    paragraph.paragraph_format.keep_with_next = keep
    if text:
        run = paragraph.add_run(text)
        set_font(run, size=size, color=color, bold=bold)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[level])
    paragraph.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}[level])
    run = paragraph.add_run(text)
    set_font(
        run,
        size={1: 16, 2: 13, 3: 12}[level],
        color=BLUE if level < 3 else DARK_BLUE,
        bold=True,
    )
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph()
    apply_numbering(paragraph, BULLET_NUM_ID)
    set_para_spacing(paragraph, after=4, line=1.25)
    run = paragraph.add_run(text)
    set_font(run, size=10.5, color=INK)
    return paragraph


def add_number(doc, text, num_id=LIST_NUM_ID):
    paragraph = doc.add_paragraph()
    apply_numbering(paragraph, num_id)
    set_para_spacing(paragraph, after=4, line=1.25)
    run = paragraph.add_run(text)
    set_font(run, size=10.5, color=INK)
    return paragraph


def add_callout(doc, title, body, fill=FILL_BLUE, tone=INK):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    set_para_spacing(paragraph, after=2)
    run = paragraph.add_run(title)
    set_font(run, size=10.5, color=tone, bold=True)
    body_paragraph = cell.add_paragraph()
    set_para_spacing(body_paragraph, after=0, line=1.2)
    body_run = body_paragraph.add_run(body)
    set_font(body_run, size=10.5, color=INK)
    add_p(doc, "", after=3)


def add_kv_table(doc, rows, widths=(2160, 7200), header=None):
    row_count = len(rows) + (1 if header else 0)
    table = doc.add_table(rows=row_count, cols=2)
    set_table_geometry(table, list(widths))
    start = 0
    if header:
        cell = table.cell(0, 0)
        cell.merge(table.cell(0, 1))
        set_cell_shading(cell, FILL_BLUE)
        paragraph = cell.paragraphs[0]
        set_para_spacing(paragraph, after=0)
        run = paragraph.add_run(header)
        set_font(run, size=10.5, color=DARK_BLUE, bold=True)
        start = 1
    for row_index, (label, value) in enumerate(rows, start=start):
        left, right = table.rows[row_index].cells
        set_cell_shading(left, FILL_GRAY)
        p_left = left.paragraphs[0]
        set_para_spacing(p_left, after=0)
        r_left = p_left.add_run(label)
        set_font(r_left, size=9.5, color=DARK_BLUE, bold=True)
        p_right = right.paragraphs[0]
        set_para_spacing(p_right, after=0, line=1.18)
        r_right = p_right.add_run(value)
        set_font(r_right, size=9.5, color=INK)
    add_p(doc, "", after=3)
    return table


def add_matrix_table(doc, headers, rows, widths, center_cols=(), font_size=8.8):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    set_table_geometry(table, widths)
    repeat_table_header(table.rows[0])
    for table_row in table.rows:
        prevent_row_split(table_row)
    for index, header in enumerate(headers):
        cell = table.cell(0, index)
        set_cell_shading(cell, FILL_BLUE)
        paragraph = cell.paragraphs[0]
        set_para_spacing(paragraph, after=0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        set_font(run, size=9, color=DARK_BLUE, bold=True)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            if column_index == 0:
                set_cell_shading(cell, "FAFBFC")
            paragraph = cell.paragraphs[0]
            set_para_spacing(paragraph, after=0, line=1.15)
            if row_index == 1 and column_index < len(row) - 1:
                paragraph.paragraph_format.keep_with_next = True
            if column_index in center_cols:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(str(value))
            set_font(run, size=font_size, color=INK, bold=(column_index == 0))
    add_p(doc, "", after=3)
    return table


def add_code_block(doc, code, title=None, font_size=7.6):
    if title:
        paragraph = add_p(doc, title, bold=True, color=DARK_BLUE, size=10.5, after=4, keep=True)
        paragraph.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, FILL_CODE)
    cell.text = ""
    for index, line in enumerate(code.rstrip().splitlines()):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        set_para_spacing(paragraph, after=0, line=1.0)
        run = paragraph.add_run(line if line else " ")
        set_font(run, name="Menlo", east_asia="PingFang SC", size=font_size, color=RGBColor(22, 27, 34))
    add_p(doc, "", after=3)


def setup_document(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    for style_name in ("Normal", "Body Text"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(11)
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25

    for index, size in ((1, 16), (2, 13), (3, 12)):
        style = styles[f"Heading {index}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE if index < 3 else DARK_BLUE
        style.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}[index])
        style.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}[index])

    add_numbering_definition(doc)

    header = section.header.paragraphs[0]
    set_para_spacing(header, after=0)
    header_run = header.add_run("智慧冷冻站 | PLC标准化与AI控制接口")
    set_font(header_run, size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_spacing(footer, after=0)
    footer_run = footer.add_run("工程受控文件 | 第 ")
    set_font(footer_run, size=9, color=MUTED)
    add_page_field(footer)
    footer_run_2 = footer.add_run(" 页")
    set_font(footer_run_2, size=9, color=MUTED)


def normalize_font_metadata(path):
    """Remove template-only font references that trigger WPS missing-font warnings."""
    temp_path = path.with_suffix(".fonts.tmp.docx")
    drawing_ns = {"a": DRAWING_NS}
    word_ns = {"w": WORD_NS}
    math_ns = {"m": MATH_NS}

    with ZipFile(path, "r") as source, ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)

            if item.filename == "word/theme/theme1.xml":
                root = etree.fromstring(data)
                for node in root.xpath(".//a:fontScheme/*/a:latin", namespaces=drawing_ns):
                    node.set("typeface", "Arial")
                for node in root.xpath(".//a:fontScheme/*/a:ea", namespaces=drawing_ns):
                    node.set("typeface", "PingFang SC")
                for node in root.xpath(".//a:fontScheme/*/a:cs", namespaces=drawing_ns):
                    node.set("typeface", "Arial")
                for node in root.xpath(".//a:fontScheme/*/a:font", namespaces=drawing_ns):
                    script = node.get("script", "")
                    node.set("typeface", "PingFang SC" if script in {"Hans", "Hant", "Jpan", "Hang"} else "Arial")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

            elif item.filename in {"word/styles.xml", "word/stylesWithEffects.xml"}:
                root = etree.fromstring(data)
                for node in root.xpath(".//w:rFonts", namespaces=word_ns):
                    for attr in ("ascii", "hAnsi", "cs"):
                        key = f"{{{WORD_NS}}}{attr}"
                        if node.get(key) not in {None, "Arial", "Menlo", "Symbol"}:
                            node.set(key, "Arial")
                    east_asia = f"{{{WORD_NS}}}eastAsia"
                    if node.get(east_asia) not in {None, "PingFang SC"}:
                        node.set(east_asia, "PingFang SC")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

            elif item.filename == "word/settings.xml":
                root = etree.fromstring(data)
                for node in root.xpath(".//m:mathFont", namespaces=math_ns):
                    node.set(f"{{{MATH_NS}}}val", "Arial")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

            elif item.filename == "word/fontTable.xml":
                root = etree.fromstring(data)
                for child in list(root):
                    root.remove(child)
                for font_name in ("Arial", "PingFang SC", "Menlo", "Symbol"):
                    font = etree.SubElement(root, f"{{{WORD_NS}}}font")
                    font.set(f"{{{WORD_NS}}}name", font_name)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

            target.writestr(item, data)

    temp_path.replace(path)


def build_doc():
    doc = Document()
    setup_document(doc)
    doc.core_properties.title = "冷站PLC标准化编程与AI系统直写实施说明"
    doc.core_properties.subject = "冷站AI控制PLC标准接口、功能块、写入边界、验收和实施工期"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "冷冻站, PLC, AI控制, OPC UA, Modbus TCP, 标准化, 安全联锁"

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(10)
    title.paragraph_format.space_after = Pt(3)
    title_run = title.add_run("冷站PLC标准化编程与AI系统直写实施说明")
    set_font(title_run, size=24, color=RGBColor(0, 0, 0), bold=True)

    subtitle = doc.add_paragraph()
    set_para_spacing(subtitle, after=14)
    subtitle_run = subtitle.add_run("适用于新建及既有冷冻站的PLC改造、AI控制接入、网关配置、现场联调和闭环验收")
    set_font(subtitle_run, size=12.5, color=MUTED)

    for label, value in (
        ("文件类型", "技术规范 / 实施SOP"),
        ("版本", "V1.0"),
        ("版本日期", date.today().isoformat()),
        ("默认控制路径", "shadow -> assisted -> enforced"),
        ("适用对象", "冷机、冷冻泵、冷却泵、冷却塔、阀门、PLC、SCADA、边缘网关"),
    ):
        paragraph = doc.add_paragraph()
        set_para_spacing(paragraph, after=2, line=1.1)
        label_run = paragraph.add_run(f"{label}: ")
        set_font(label_run, size=10.5, color=INK, bold=True)
        value_run = paragraph.add_run(value)
        set_font(value_run, size=10.5, color=INK)

    rule = doc.add_paragraph()
    set_para_spacing(rule, before=7, after=12)
    paragraph_border_bottom(rule)

    add_callout(
        doc,
        "核心结论",
        "PLC端可以做到约85%-95%标准化。系统可直接写入独立的AI命令区，但不得直接写设备输出、原始PID输出、保护上下限和急停联锁；PLC程序下载仍必须经过备份、编译、在线比对和人工授权。新项目应预装标准AI桥，既有项目建议先增加最小L1桥接层。",
        fill=FILL_WARN,
        tone=RISK,
    )

    add_heading(doc, "文档导航", 1)
    for item in (
        "第1-2章：明确系统直写边界和推荐架构。",
        "第3-6章：给出PLC数据契约、标准功能块、设备控制逻辑和PLC伪代码。",
        "第7-8章：说明自动生成能力、不同PLC平台适配和B25实施建议。",
        "第9-10章：规定FAT/SAT验收、责任边界、工期和交付物。",
        "附录：参数基线、拒绝原因码、现场交付检查表。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "1. 目标、范围与边界", 1)
    add_p(
        doc,
        "本说明用于把冷站AI优化系统与不同品牌PLC之间的接口收敛成统一工程标准。目标是减少每个项目重复编程和现场工作量，同时保证PLC始终拥有安全联锁、顺控、PID、限幅和回退的最终决定权。",
    )

    add_heading(doc, "1.1 需要区分的三类写入", 2)
    add_matrix_table(
        doc,
        ["写入类型", "系统是否可自动执行", "工程边界"],
        [
            ("运行目标值", "可以", "只写AI命令区中的目标、Trim、台数请求或模式申请。"),
            ("控制参数", "受限", "策略上下限可在系统侧配置，但PLC保护极限、PID参数和联锁阈值不得自动修改。"),
            ("PLC程序下载", "不自动", "系统可生成厂家导入包；下载必须由工程软件完成并经过人工变更审批。"),
        ],
        [2100, 2100, 5160],
        center_cols=(1,),
    )

    add_heading(doc, "1.2 系统直写允许与禁止矩阵", 2)
    add_matrix_table(
        doc,
        ["对象", "结论", "推荐实现"],
        [
            ("冷冻水供水温度目标", "允许", "写AI_ChwsTarget，PLC限幅和斜率执行。"),
            ("冷冻泵/冷却泵频率", "仅允许Trim", "写系统级修正量，由PLC叠加本地PID并分配给运行泵。"),
            ("冷却水温度/Approach", "允许", "写AI_TcwsTarget或AI_ApproachTarget，PLC结合湿球和下限二次校验。"),
            ("冷机台数", "允许请求", "写ChillerCountRequest；PLC完成选机、阀门、流量证明和启停顺控。"),
            ("冷机/泵/塔直接启停", "首期禁止", "必须通过PLC设备适配块和顺控状态机。"),
            ("手自动模式", "高风险", "只写模式申请，不直接覆盖现场手动或检修挂牌。"),
            ("PID输出、变频器输出", "禁止", "保留在PLC内部。"),
            ("急停、联锁、保护极限", "禁止", "任何系统和AI都不能覆盖。"),
        ],
        [2400, 1500, 5460],
        center_cols=(1,),
    )

    add_callout(
        doc,
        "验收真值",
        "HTTP 200、网关accepted、PLC寄存器写入成功都不等于设备完成执行。最终真值必须至少包括：PLC ACK、ActiveSeq一致、ActiveValue回读、设备实际反馈达到目标、联锁未触发，以及必要时的自动回退记录。",
        fill=FILL_RED,
        tone=RISK,
    )

    add_heading(doc, "2. 推荐总体架构", 1)
    add_code_block(
        doc,
        """云端/本地AI优化平台
  -> 语义命令、审批、命令编号、TTL、回退值
边缘PLC执行网关
  -> 鉴权、白名单、点位映射、缓存、防重复、审计
DB_AI_COMMAND_A/B
  -> 双缓冲写入，CommitSeq最后提交
FB_AI_CommandGate + FB_AI_SafetyGate
  -> 校验、仲裁、ACK/REJECT
FB_AI_TargetManager + 设备适配块
  -> 限幅、斜率、死区、顺控、PID
现场设备
  -> 冷机、冷冻泵、冷却泵、冷却塔、阀门
DB_AI_STATUS
  -> 生效值、设备反馈、联锁、回退和拒绝原因""",
        "控制链路",
        font_size=8.3,
    )

    add_heading(doc, "2.1 职责划分", 2)
    add_matrix_table(
        doc,
        ["层级", "负责", "不得承担"],
        [
            ("AI平台", "负荷预测、COP寻优、目标值建议、策略版本、审批和效果评估。", "秒级保护、直接设备输出、绕过PLC联锁。"),
            ("执行网关", "协议适配、命令队列、TTL、白名单、双缓冲写入、回执和审计。", "自行决定安全边界或任意写点。"),
            ("PLC标准桥", "完整性校验、模式仲裁、限幅、斜率、超时、回退、状态反馈。", "复杂经济优化和模型训练。"),
            ("原有PLC控制", "设备保护、联锁、顺控、PID、最小启停、故障处理。", "接受未经校验的外部值。"),
            ("SCADA", "显示、人工操作、报警、趋势和权限。", "把通信成功显示成设备执行成功。"),
        ],
        [1800, 3780, 3780],
    )

    add_heading(doc, "2.2 控制权限优先级", 2)
    add_number(doc, "PLC急停、设备保护和硬联锁，任何外部命令都不能覆盖。")
    add_number(doc, "现场手动、检修、挂牌和消防联动优先于远程及AI控制。")
    add_number(doc, "本地人工操作优先于云端人工和AI命令。")
    add_number(doc, "AI命令只能在AI_ENABLE、自动模式、安全输入和数据质量全部有效时被PLC接受。")
    add_number(doc, "通信异常或命令过期时，PLC自动进入local_fallback，不依赖AI平台恢复。")

    doc.add_page_break()
    add_heading(doc, "3. 标准PLC数据契约", 1)
    add_p(
        doc,
        "所有项目应建立独立的命令区和状态区。OPC UA优先使用结构化Tag；Modbus TCP使用固定寄存器块，并采用双缓冲、校验值和最后提交字段，避免多寄存器写入过程中产生半条命令。",
    )

    add_heading(doc, "3.1 DB_AI_COMMAND命令区", 2)
    add_matrix_table(
        doc,
        ["字段", "类型", "含义", "要求"],
        [
            ("ProtocolVersion", "UINT", "接口版本", "不兼容版本必须拒绝。"),
            ("CommandSeq", "UDINT", "命令序号", "单调递增，用于防重复。"),
            ("TimestampUtc", "DT/LINT", "命令时间", "用于判断时钟和过期。"),
            ("TTLSeconds", "UINT", "有效时间", "到期未执行必须拒绝。"),
            ("CommandSource", "USINT", "本地人工/云端人工/本地AI/云端AI", "参与权限仲裁。"),
            ("RequestedMode", "USINT", "off/shadow/assisted/enforced", "PLC可降级，不能被系统强制升级。"),
            ("TargetMask", "DWORD", "本次包含哪些目标", "未置位字段不得改变当前值。"),
            ("TargetValues", "STRUCT", "温度、Trim、台数、模式请求", "统一工程单位，不传原始比例值。"),
            ("RollbackValues", "STRUCT", "回退值", "assisted/enforced必须提供。"),
            ("PayloadChecksum", "UDINT", "载荷校验", "Modbus顺序写入时强制使用。"),
            ("CommitSeq", "UDINT", "提交标志", "必须最后写入，并等于CommandSeq。"),
        ],
        [2200, 1350, 3040, 2770],
        center_cols=(1,),
        font_size=8.4,
    )

    add_heading(doc, "3.2 DB_AI_STATUS状态区", 2)
    add_matrix_table(
        doc,
        ["字段", "用途", "系统验收条件"],
        [
            ("LastReceivedSeq", "PLC已完整收到的命令", "等于CommandSeq。"),
            ("LastAcceptedSeq", "PLC接受的命令", "拒绝时保持旧值。"),
            ("CommandState", "IDLE/VALIDATING/ACCEPTED/APPLYING/ACTIVE/REJECTED/ROLLING_BACK", "不能只停在ACCEPTED。"),
            ("RejectReason", "标准拒绝原因码", "必须可追溯到联锁或参数。"),
            ("SafetyMask", "当前安全条件位图", "关键位全部为真。"),
            ("ActiveValues", "PLC实际采用目标", "与请求值或限幅后值一致。"),
            ("ActualFeedback", "设备频率、温度、状态、台数", "在允许偏差和超时内达到目标。"),
            ("FallbackActive", "本地回退已激活", "回退时必须同步记录。"),
        ],
        [2300, 3200, 3860],
    )

    add_heading(doc, "3.3 原子命令生命周期", 2)
    for step in (
        "系统选择非活动缓冲区A或B，先写ProtocolVersion、CommandSeq、目标值、回退值和TTL。",
        "系统计算PayloadChecksum，最后写CommitSeq。",
        "PLC仅在CommitSeq=CommandSeq且校验正确时复制到内部工作区。",
        "PLC依次检查版本、时间、权限、手自动、故障、数据质量、上下限和设备可用性。",
        "PLC返回ACCEPTED或REJECTED；接受后按斜率和顺控进入APPLYING。",
        "设备反馈满足条件后进入ACTIVE；若超时或性能恶化则进入ROLLING_BACK。",
        "系统回读LastAcceptedSeq、ActiveValues和ActualFeedback后，才标记命令VERIFIED。",
    ):
        add_number(doc, step, LIST_COMMAND_NUM_ID)

    add_heading(doc, "4. PLC端标准编程模块", 1)
    add_matrix_table(
        doc,
        ["模块", "标准功能", "项目特定配置"],
        [
            ("FB_AI_CommandGate", "版本、序号、时间戳、TTL、校验、防重复、ACK/REJECT。", "通信周期和时钟来源。"),
            ("FB_AI_Arbitration", "现场手动、本地人工、远程人工、AI优先级。", "现有模式位和检修挂牌信号。"),
            ("FB_AI_SafetyGate", "安全条件位图、故障闭锁、数据质量闭锁。", "流量、压力、温度、阀位、故障点映射。"),
            ("FB_AI_TargetManager", "限幅、死区、斜率、最小保持、变化率检查。", "各站参数上下限。"),
            ("FB_AI_Watchdog", "心跳、通信超时、本地回退。", "超时时间和回退模式。"),
            ("FB_AI_Rollback", "回退到0Hz Trim、原设定值或最近稳定值。", "回退值来源和闭锁时长。"),
            ("FB_AI_ChillerAdapter", "台数/目标温度请求转为原有冷机顺控输入。", "冷机编号、优先级、最小启停时间。"),
            ("FB_AI_PumpAdapter", "系统级Trim叠加本地PID并分配给运行泵。", "泵组、最低频率、最小流量。"),
            ("FB_AI_TowerAdapter", "Tcws/Approach目标转为风机和加减塔请求。", "湿球点、最低频率、加减塔延时。"),
            ("DB_AI_Audit", "保存最近命令、拒绝、回退和关键快照。", "保留条数和SCADA上传方式。"),
        ],
        [2460, 3900, 3000],
        font_size=8.5,
    )

    add_callout(
        doc,
        "标准化原则",
        "标准库不直接替换原有设备控制程序。它位于AI系统与原有顺控/PID之间，把外部请求转换成经过安全校验的内部目标。现场只需映射设备点、联锁信号、范围和回退值。",
        fill=FILL_GOOD,
        tone=GOOD,
    )

    doc.add_page_break()
    add_heading(doc, "5. 设备控制逻辑要求", 1)
    add_heading(doc, "5.1 控制目标", 2)
    add_bullet(doc, "在满足供冷、末端舒适和设备安全的前提下，提高系统综合COP并降低总功率。")
    add_bullet(doc, "避免冷机、泵和冷却塔频繁启停，避免温度锯齿波和设备在低效区运行。")
    add_bullet(doc, "AI只给出目标值、修正量或台数请求，PLC负责最终执行、拒绝和回退。")

    add_heading(doc, "5.2 输入变量", 2)
    add_matrix_table(
        doc,
        ["对象", "必要输入", "数据质量要求"],
        [
            ("系统", "冷量、总功率、系统COP、负荷率、模式、告警。", "时间同步；关键功率和温度不得坏点。"),
            ("冷机", "运行/故障、负载率、功率、蒸发/冷凝温压、最小启停计时。", "状态和功率反馈一致。"),
            ("冷冻泵", "供回水温、Delta-T、流量、压差、频率、末端阀位和室温风险。", "最小流量、频率反馈必须有效。"),
            ("冷却泵", "冷却水温差、流量、频率、冷凝压力/温度。", "冷凝器风险输入不得缺失。"),
            ("冷却塔", "室外湿球、冷却水供回水温、Approach、风机频率、运行台数、温度变化率。", "湿球和供水温必须在合理范围。"),
            ("阀门", "开度、到位、故障、联锁、动作超时。", "指令与反馈偏差可检测。"),
        ],
        [1800, 4560, 3000],
        font_size=8.5,
    )

    add_heading(doc, "5.3 输出控制量", 2)
    add_matrix_table(
        doc,
        ["控制对象", "AI输出", "PLC执行方式"],
        [
            ("冷机", "ChillerCountRequest、AI_ChwsTarget", "PLC选择机组并执行阀门、泵、流量证明和启停顺控。"),
            ("冷冻泵", "AI_ChwpFreqTrim_Hz或AI_PumpDpTarget", "叠加本地PID、限幅、斜率和最小流量保护。"),
            ("冷却泵", "AI_CwpFreqTrim_Hz", "结合冷凝压力和冷却水温保护。"),
            ("冷却塔", "AI_TcwsTarget或AI_ApproachTarget", "根据湿球、最低Approach、台数和风机频率执行。"),
            ("阀门", "ValvePositionRequest或SequenceRequest", "由PLC进行到位确认、超时和联锁。"),
        ],
        [1900, 3000, 4460],
    )

    add_heading(doc, "5.4 防震荡机制", 2)
    add_matrix_table(
        doc,
        ["机制", "初始建议", "目的"],
        [
            ("优化周期", "5分钟", "避免AI追随短时噪声。"),
            ("泵频单步", "不超过1Hz", "防止压差和流量突变。"),
            ("Tcws单步", "不超过0.5摄氏度", "防止冷凝工况和风机功率剧烈变化。"),
            ("温度死区", "0.2-0.5摄氏度", "抑制设定值来回切换。"),
            ("最小保持时间", "5-10分钟", "保证效果可观察。"),
            ("回退后闭锁", "15分钟", "避免异常后立即再次优化。"),
            ("冷机最小运行/停机", "30-60/15-30分钟", "保护压缩机并减少频繁启停。"),
            ("冷却塔加减塔延时", "10-20分钟", "避免台数切换震荡。"),
        ],
        [2500, 2100, 4760],
        center_cols=(1,),
    )

    add_heading(doc, "5.5 异常保护与回退", 2)
    for item in (
        "急停、设备跳闸、检修、现场手动、消防联动：立即拒绝外部命令并保持本地安全控制。",
        "心跳超时、时钟异常、命令过期、Checksum错误：不得使用新目标，进入local_fallback。",
        "温度、流量、压力、功率或状态质量无效：冻结优化并使用最近稳定目标。",
        "末端压差过低、阀位大面积接近全开、室温超限：禁止冷冻泵继续降频。",
        "冷凝压力升高、冷却水温超限、最小流量不足：禁止冷却泵和冷却塔进一步降频。",
        "执行后系统COP恶化或总功率上升超过阈值：按斜率回退并记录原因。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "6. PLC可实现伪代码", 1)
    add_code_block(
        doc,
        """// FB_AI_CommandGate - IEC 61131-3 Structured Text示意
NewCommit := (Cmd.CommitSeq = Cmd.CommandSeq)
             AND (Cmd.CommandSeq <> Status.LastReceivedSeq);

IF EmergencyAlarm OR SafetyTrip THEN
    Reject(REJECT_SAFETY_TRIP);
    UseLocalSafeControl();
ELSIF LocalManual OR MaintenanceMode THEN
    Reject(REJECT_LOCAL_MANUAL);
    UseLocalControl();
ELSIF HeartbeatTimeout > T#60s THEN
    Status.FallbackActive := TRUE;
    StartRollback(REJECT_HEARTBEAT_TIMEOUT);
ELSIF NewCommit THEN
    Status.LastReceivedSeq := Cmd.CommandSeq;
    Status.CommandState := VALIDATING;

    IF Cmd.ProtocolVersion <> SUPPORTED_VERSION THEN
        Reject(REJECT_VERSION);
    ELSIF NOT ChecksumOK(Cmd) THEN
        Reject(REJECT_CHECKSUM);
    ELSIF CommandExpired(Cmd.TimestampUtc, Cmd.TTLSeconds) THEN
        Reject(REJECT_EXPIRED);
    ELSIF NOT AiEnable OR NOT SafetyInputsOK THEN
        Reject(REJECT_SAFETY_INPUT);
    ELSIF NOT TargetsInEngineeringRange(Cmd.TargetValues) THEN
        Reject(REJECT_OUT_OF_RANGE);
    ELSE
        WorkTarget := ClampToOperatingLimits(Cmd.TargetValues);
        WorkRollback := ValidateRollback(Cmd.RollbackValues);
        Status.LastAcceptedSeq := Cmd.CommandSeq;
        Status.CommandState := ACCEPTED;
        StartApply := TRUE;
    END_IF;
END_IF;""",
        "命令门禁",
        font_size=7.4,
    )

    doc.add_page_break()
    add_code_block(
        doc,
        """// FB_AI_TargetManager - 目标执行、反馈与回退
IF StartApply THEN
    ActiveTarget := RampLimit(
        Current := ActiveTarget,
        Requested := WorkTarget,
        MaxStep := Config.MaxStepPerCycle
    );

    IF Abs(ActiveTarget - CurrentLocalTarget) < Config.Deadband THEN
        HoldCurrentTarget();
    ELSE
        ApplyToLocalController(ActiveTarget);
    END_IF;

    Status.CommandState := APPLYING;
END_IF;

IF FeedbackReached(ActiveTarget, ActualFeedback, Config.Tolerance) THEN
    Status.ActiveValues := ActiveTarget;
    Status.CommandState := ACTIVE;
ELSIF ApplyTimeout OR PerformanceDegraded OR NewSafetyTrip THEN
    Status.CommandState := ROLLING_BACK;
    ActiveTarget := RampLimit(ActiveTarget, WorkRollback, Config.RollbackStep);
    LockAiControl(Config.RollbackLockout);
END_IF;""",
        "目标管理与回退",
        font_size=7.4,
    )

    add_code_block(
        doc,
        """// 泵侧建议：Trim而不是绝对频率
ChwpTargetHz := LocalChwpPidHz + AI_ChwpFreqTrim_Hz;
ChwpTargetHz := LIMIT(ChwpMinHz, ChwpTargetHz, ChwpMaxHz);

IF NOT MinimumChilledFlowOK OR TerminalPressureLow OR ValveRiskHigh THEN
    AI_ChwpFreqTrim_Hz := 0.0;
END_IF;

// 塔侧建议：结合湿球和最低冷凝器进水温
MinTcwsByWetBulb := WetBulbC + MinApproachC;
SafeTcwsTarget := MAX(AI_TcwsTarget, MinTcwsByWetBulb, MinCondenserInletTempC);
SafeTcwsTarget := RampLimit(CurrentTcwsTarget, SafeTcwsTarget, 0.5);""",
        "泵侧与塔侧关键逻辑",
        font_size=7.4,
    )

    doc.add_page_break()
    add_heading(doc, "7. 标准化、代码生成与系统直写", 1)
    add_heading(doc, "7.1 三级接入模式", 2)
    add_matrix_table(
        doc,
        ["等级", "PLC改造", "允许能力", "适用场景"],
        [
            ("L0兼容接入", "不改PLC，使用原有点", "只读、shadow；必要时受控人工写入", "无法停机或暂时无PLC源程序。"),
            ("L1标准桥接", "增加AI命令区、门禁、回退和状态区", "人工审批assisted", "既有项目推荐。"),
            ("L2完整标准", "增加设备适配块和完整状态机", "验收后可考虑enforced", "新建项目和批量复制项目。"),
        ],
        [1600, 3300, 2260, 2200],
        center_cols=(0,),
    )

    add_heading(doc, "7.2 系统可自动生成的交付物", 2)
    for item in (
        "统一点位字典、工程单位、比例、读写属性、权限等级和写入风险清单。",
        "DB_AI_COMMAND、DB_AI_STATUS、UDT/DUT/标签和标准功能块导入包。",
        "OPC UA Tag映射、Modbus地址表、字节序和缩放配置。",
        "网关运行配置、语义目标映射、写入白名单、范围和TTL。",
        "FAT测试用例、模拟器数据、拒绝原因测试、回退测试和SAT检查表。",
        "站点runtime-config、shadow/assisted门禁和验收报告模板。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "7.3 不能取消的人工步骤", 2)
    for item in (
        "读取并备份现场PLC原程序，确认PLC型号、固件、冗余和在线变更限制。",
        "核对原有联锁、手自动语义、PID输出位置、设备顺控和故障处理。",
        "在厂家工程软件中编译、离线仿真、在线比对并下载标准块。",
        "现场逐项验证写入、ACK、ActiveValue、设备反馈、通信超时和回退。",
        "由甲方、自控集成商和系统实施方共同签署assisted/enforced放行记录。",
    ):
        add_number(doc, item, LIST_MANUAL_NUM_ID)

    add_heading(doc, "7.4 不同PLC平台标准包", 2)
    add_matrix_table(
        doc,
        ["PLC平台", "标准包形式", "推荐通信"],
        [
            ("西门子S7-1200/1500", "SCL源文件、UDT、FB、实例DB、全局DB", "OPC UA或S7协议"),
            ("CODESYS", "DUT、FB Library、GVL", "OPC UA或Modbus TCP"),
            ("罗克韦尔", "UDT、AOI、L5X导入包", "EtherNet/IP或OPC UA"),
            ("三菱", "全局标签、FB库、结构体", "MC协议或OPC UA网关"),
            ("施耐德", "DDT、DFB、变量表", "OPC UA或Modbus TCP"),
        ],
        [2500, 4060, 2800],
    )

    add_callout(
        doc,
        "一键部署的准确含义",
        "系统可以一键生成PLC适配包、点位配置、网关配置和测试清单；自控工程师首次导入并验收后，AI系统可以直接、标准化地写入目标命令。所谓一键部署不应包含未经人工确认的PLC在线下载。",
        fill=FILL_GOOD,
        tone=GOOD,
    )

    add_heading(doc, "8. B25项目实施建议", 1)
    add_p(
        doc,
        "B25当前可写的冷却回水手动值不是独立AI目标点；B25_AI_ChwpFreqTrim_Hz和B25_AI_CwpFreqTrim_Hz目前也只是系统shadow映射名称，不代表现场PLC已存在真实点。因此当前边界仍应保持GO_SHADOW/GO_SHADOW_ONLY，不宜直接开启真实闭环。",
    )

    add_heading(doc, "8.1 推荐新增的PLC接口", 2)
    add_matrix_table(
        doc,
        ["点/结构", "方向", "用途"],
        [
            ("B25_AI_COMMAND_A/B", "系统 -> PLC", "完整命令双缓冲区。"),
            ("B25_AI_STATUS", "PLC -> 系统", "ACK、拒绝、ActiveValue、联锁和回退。"),
            ("B25_AI_ENABLE", "现场/PLC", "AI控制总许可。"),
            ("B25_AI_TcwsTarget", "系统 -> PLC", "冷却水供水温度目标。"),
            ("B25_AI_ChwpFreqTrim", "系统 -> PLC", "冷冻泵组频率修正量。"),
            ("B25_AI_CwpFreqTrim", "系统 -> PLC", "冷却泵组频率修正量。"),
            ("B25_AI_ROLLBACK_REQ", "系统/PLC", "人工或自动回退申请。"),
            ("B25_AI_SAFETY_MASK", "PLC -> 系统", "最小流量、末端风险、冷凝风险、手自动等安全状态。"),
        ],
        [3100, 1800, 4460],
        center_cols=(1,),
    )

    add_heading(doc, "8.2 B25第一阶段控制范围", 2)
    add_bullet(doc, "冷却塔：只接入AI_TcwsTarget，不直接写冷却塔温度手自动点。")
    add_bullet(doc, "冷冻泵和冷却泵：只接入组级Trim，第一版不做泵台数启停。")
    add_bullet(doc, "冷机：只读和组合建议，第一阶段不开放冷机自动启停。")
    add_bullet(doc, "所有命令先shadow，再进入人工审批assisted；enforced需要另行SAT和书面放行。")

    add_heading(doc, "8.3 B25实施步骤", 2)
    for step in (
        "导出PLC源程序和最终点位表，核对现有41413手动值、频率修正点和手自动语义。",
        "导入L1标准桥：命令区、状态区、CommandGate、SafetyGate、TargetManager和Rollback。",
        "把新AI目标映射到原有本地PID/顺控输入，不直接映射到变频器输出或设备启停线圈。",
        "离线模拟正常命令、越限、重复、过期、手动、故障、心跳超时和回退。",
        "现场先只读和shadow，核对系统建议与PLC可接受范围。",
        "人工审批assisted小步测试：泵频每步不超过1Hz，Tcws每步不超过0.5摄氏度。",
        "完成回读、设备反馈、COP不劣化和自动回退验收后，再评估enforced。",
    ):
        add_number(doc, step, LIST_B25_NUM_ID)

    doc.add_page_break()
    add_heading(doc, "9. FAT、SAT与闭环验收", 1)
    add_heading(doc, "9.1 FAT离线测试", 2)
    add_matrix_table(
        doc,
        ["测试项", "输入", "预期结果"],
        [
            ("正常命令", "合法目标、有效TTL、安全条件正常", "ACCEPTED -> APPLYING -> ACTIVE。"),
            ("重复命令", "相同CommandSeq再次提交", "不重复执行，状态保持。"),
            ("半条命令", "CommitSeq不一致或Checksum错误", "拒绝，不改变ActiveValue。"),
            ("过期命令", "时间戳超过TTL", "REJECT_EXPIRED。"),
            ("越限目标", "超过工程范围", "REJECT_OUT_OF_RANGE。"),
            ("现场手动", "LocalManual=TRUE", "拒绝AI并保持现场控制。"),
            ("安全故障", "最小流量/冷凝风险/设备故障", "拒绝或自动回退。"),
            ("心跳丢失", "超过60秒无心跳", "local_fallback，按策略回退。"),
            ("反馈超时", "目标未达到或状态不一致", "ROLLING_BACK并闭锁。"),
        ],
        [2300, 3500, 3560],
        font_size=8.5,
    )

    add_heading(doc, "9.2 SAT现场测试", 2)
    for item in (
        "在线程序与备份版本一致，标准块下载范围清楚，没有无关程序变化。",
        "OPC UA/Modbus写入单位、比例、字节序、浮点格式和地址正确。",
        "PLC ACK与ActiveSeq一致，ActiveValue可回读，SCADA显示真实状态。",
        "手动、检修、故障、通信中断和现场优先级均能阻断AI命令。",
        "泵频、温度目标、阀门和设备状态按照限幅与斜率变化，无突跳。",
        "回退后设备回到安全目标，AI闭锁时间生效，审计记录完整。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "9.3 放行条件", 2)
    add_matrix_table(
        doc,
        ["阶段", "允许动作", "放行证据"],
        [
            ("shadow", "只生成拟执行命令，不写PLC", "点位映射、范围、建议值和风险审阅通过。"),
            ("assisted", "人工审批后小步写入", "FAT/SAT通过，回读、联锁、回退和审计完整。"),
            ("enforced", "门禁满足后自动执行", "典型工况覆盖、效果不劣化、异常回退验证和书面签字。"),
        ],
        [1800, 3000, 4560],
    )

    add_callout(
        doc,
        "禁止误判",
        "不能用网页可打开、BFF ok=true、网关accepted、寄存器有数值或单次运行成功代替闭环验收。真实闭环必须闭合到PLC生效值、设备反馈、性能结果、回退记录和现场签字。",
        fill=FILL_RED,
        tone=RISK,
    )

    add_heading(doc, "10. 工期、职责与交付物", 1)
    add_heading(doc, "10.1 预计工期", 2)
    add_matrix_table(
        doc,
        ["场景", "PLC/系统配置", "现场验证", "预计总工期"],
        [
            ("首个PLC品牌标准库", "3-5天", "1-2天FAT", "4-7天"),
            ("B25增加L1标准桥", "0.5-1天", "1-2天SAT", "约2-3天"),
            ("后续同品牌同架构项目", "2-4小时", "0.5-1天SAT", "最快1天"),
            ("完整无人值守闭环", "在L2基础上", "覆盖典型工况", "通常3-10个工作日"),
        ],
        [2700, 2100, 2100, 2460],
        center_cols=(1, 2, 3),
    )

    add_heading(doc, "10.2 责任边界", 2)
    add_matrix_table(
        doc,
        ["责任方", "主要工作", "签字责任"],
        [
            ("AI系统方", "语义命令、算法边界、网关配置、审计、效果评估和回退触发。", "系统功能和命令契约。"),
            ("PLC自控方", "导入标准块、现场映射、联锁、顺控、PID、程序备份和下载。", "PLC安全逻辑和程序变更。"),
            ("设备厂家", "冷机允许条件、通讯点、保护边界和启停约束。", "设备控制条件。"),
            ("甲方运维", "运行模式、检修挂牌、现场操作、验收窗口和放行权限。", "现场运行和控制模式。"),
            ("项目经理", "版本、风险、测试记录、变更审批和最终交付。", "交付完整性。"),
        ],
        [1800, 4860, 2700],
    )

    add_heading(doc, "10.3 标准交付物", 2)
    for item in (
        "PLC标准库及版本说明、厂家导入包、源程序备份和在线差异报告。",
        "AI命令区/状态区数据字典、OPC UA/Modbus映射表和写入白名单。",
        "站点参数表、上下限、斜率、死区、TTL、回退值和拒绝原因码。",
        "FAT记录、SAT记录、shadow报告、assisted执行记录和回退测试证据。",
        "控制权限矩阵、变更审批、甲方/自控/系统三方签字和最终竣工版本。",
    ):
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "附录A. 推荐参数基线", 1)
    add_matrix_table(
        doc,
        ["参数", "初始建议", "现场调整依据"],
        [
            ("AI心跳周期", "5秒", "网络和PLC扫描周期。"),
            ("心跳超时", "60秒", "不应依赖云端恢复。"),
            ("命令TTL", "300秒", "优化周期和审批耗时。"),
            ("执行反馈超时", "10-30秒", "设备响应速度。"),
            ("泵频Trim范围", "-5至+3Hz", "最低流量、压差和变频器范围。"),
            ("泵频单步", "不超过1Hz", "流量和末端稳定性。"),
            ("Tcws单步", "不超过0.5摄氏度", "冷凝器和塔侧稳定性。"),
            ("温度死区", "0.2-0.5摄氏度", "传感器噪声和控制周期。"),
            ("最小保持时间", "5-10分钟", "观察功率和温度效果。"),
            ("回退闭锁", "15分钟", "异常后稳定恢复。"),
            ("冷机最小运行", "30-60分钟", "厂家要求和负荷变化。"),
            ("冷机最小停机", "15-30分钟", "厂家要求。"),
            ("冷却塔加减塔延时", "10-20分钟", "水温变化率和负荷。"),
        ],
        [3300, 2300, 3760],
        center_cols=(1,),
    )

    add_heading(doc, "附录B. 标准拒绝原因码", 1)
    add_matrix_table(
        doc,
        ["代码", "名称", "含义"],
        [
            ("0", "NONE", "无拒绝。"),
            ("1", "AI_DISABLED", "AI_ENABLE未打开。"),
            ("2", "LOCAL_MANUAL", "现场手动或检修挂牌。"),
            ("3", "SAFETY_TRIP", "急停、设备故障或安全联锁。"),
            ("4", "COMMAND_EXPIRED", "命令超过TTL。"),
            ("5", "DUPLICATE_COMMAND", "CommandSeq重复。"),
            ("6", "CHECKSUM_ERROR", "命令数据不完整或校验失败。"),
            ("7", "TARGET_OUT_OF_RANGE", "目标超出工程范围。"),
            ("8", "DATA_QUALITY_BAD", "关键输入质量无效。"),
            ("9", "DEVICE_UNAVAILABLE", "目标设备不可用或顺控条件不满足。"),
            ("10", "FEEDBACK_TIMEOUT", "执行后反馈未达到。"),
            ("11", "PERFORMANCE_DEGRADED", "总功率/COP或关键工况恶化。"),
            ("12", "ROLLBACK_LOCKOUT", "回退闭锁期间拒绝新命令。"),
            ("13", "PROTOCOL_VERSION", "接口版本不兼容。"),
        ],
        [1100, 3000, 5260],
        center_cols=(0,),
    )

    add_heading(doc, "附录C. 现场交付检查表", 1)
    for item in (
        "已取得PLC最终源程序、CPU型号、固件版本、网络拓扑和程序备份。",
        "点位表包含Tag、地址、类型、读写属性、单位、比例、权限和报警等级。",
        "AI命令区和状态区已与原控制区隔离，并配置写入白名单。",
        "所有AI目标都映射到本地目标/请求，不直接映射到输出线圈或变频器输出。",
        "手动、检修、故障、最小流量、冷凝风险和末端风险均可闭锁AI命令。",
        "重复、过期、半条命令、越限、心跳丢失和反馈超时测试全部通过。",
        "PLC ACK、ActiveSeq、ActiveValue和设备实际反馈可在SCADA及系统中审计。",
        "回退值、回退斜率、回退闭锁和本地fallback经过现场验证。",
        "shadow、assisted、enforced权限分开，enforced默认关闭。",
        "甲方、PLC自控方、AI系统方完成版本和放行签字。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "附录D. 项目依据文件", 1)
    add_kv_table(
        doc,
        [
            ("总体架构", "docs/CLOUD_LOCAL_GATEWAY_CONTROL_ARCHITECTURE_V1.md"),
            ("B25塔侧", "docs/B25_TOWER_APPROACH_POINT_MAPPING_CURRENT.md"),
            ("泵侧映射", "docs/PUMP_DELTA_T_CONTROL_MAPPING_TEMPLATE_CURRENT.md"),
            ("B25点位接入", "docs/field-data/optimize-demo-140/B25_POINT_TABLE_ONLY_ONBOARDING_CURRENT.md"),
        ],
        header="当前项目内部依据",
    )

    add_callout(
        doc,
        "最终原则",
        "系统负责优化、审批、命令和审计；PLC负责执行、安全、联锁和回退。标准化的目标不是减少安全步骤，而是把安全步骤固化成每个项目都必须通过的标准程序和验收证据。",
        fill=FILL_BLUE,
        tone=DARK_BLUE,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    normalize_font_metadata(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT.resolve())
