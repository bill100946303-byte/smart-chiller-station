from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/deployment/新项目目标Mac_Codex一键安装部署手册.docx")


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(28, 35, 43)
MUTED = RGBColor(90, 98, 110)
RISK = RGBColor(155, 28, 28)
GOOD = RGBColor(24, 112, 72)
FILL_BLUE = "E8EEF5"
FILL_GRAY = "F2F4F7"
FILL_CODE = "F7F9FB"
FILL_WARN = "FFF4E5"
FILL_GOOD = "EAF7F0"
BORDER = "D0D7DE"


INSTALL_SCRIPT = r'''#!/usr/bin/env bash
set -euo pipefail

# ===== 必填/可改变量 =====
# 目标 Codex 执行前必须把 REPO_URL 换成真实仓库地址。
REPO_URL="${REPO_URL:-}"
DEPLOY_DIR="${DEPLOY_DIR:-$HOME/Documents/智慧冷冻站}"
SITE_ID="${SITE_ID:-new_chiller_site_001}"
SITE_NAME="${SITE_NAME:-新冷站项目}"

# 默认安全模式：只读 + shadow，不真实写 PLC。
BFF_PORT="${BFF_PORT:-8787}"
FRONTEND_PORT="${FRONTEND_PORT:-3001}"
ADMIN_PORT="${ADMIN_PORT:-3002}"
PLC_GATEWAY_PORT="${PLC_GATEWAY_PORT:-8877}"
APP_MODE="${APP_MODE:-site-local}"
APP_MODE_LABEL="${APP_MODE_LABEL:-现场本地联调}"
READ_ONLY_MODE="${READ_ONLY_MODE:-1}"
PLC_GATEWAY_MODE="${PLC_GATEWAY_MODE:-mock}"   # mock | real | none
ENABLE_REAL_PLC_WRITE="${ENABLE_REAL_PLC_WRITE:-0}"
I_UNDERSTAND_REAL_PLC_WRITE="${I_UNDERSTAND_REAL_PLC_WRITE:-NO}"

LEGACY_BASE_URL="${LEGACY_BASE_URL:-https://www.ssge.com.cn:8098}"
REALTIME_PARAMS_BASE_URL="${REALTIME_PARAMS_BASE_URL:-https://ln.szgreenenergy.com}"
REALTIME_PARAMS_TIMEOUT_MS="${REALTIME_PARAMS_TIMEOUT_MS:-1500}"
MIN_CONDENSER_INLET_TEMP_C="${MIN_CONDENSER_INLET_TEMP_C:-22}"
ADMIN_DB_FILE="${ADMIN_DB_FILE:-$DEPLOY_DIR/apps/chiller-bff/.local/admin.sqlite}"

say() { printf '\n[%s] %s\n' "$(date '+%H:%M:%S')" "$*"; }
fail() { printf '\n[FAIL] %s\n' "$*" >&2; exit 1; }
warn() { printf '\n[WARN] %s\n' "$*" >&2; }

if [ -z "$REPO_URL" ]; then
  fail "REPO_URL 未设置。示例：REPO_URL=https://github.com/org/repo.git bash install_chiller_target_mac.sh"
fi

say "1/9 检查并准备 Homebrew / Git / Node 22 / nginx"
if ! command -v brew >/dev/null 2>&1; then
  say "Homebrew 未安装，开始安装。若目标机策略禁止自动安装，请先人工安装后重跑。"
  NONINTERACTIVE=1 /bin/bash -c "$(curl -fsSL https://raw.githubusercontent.com/Homebrew/install/HEAD/install.sh)"
fi
if [ -x /opt/homebrew/bin/brew ]; then
  eval "$(/opt/homebrew/bin/brew shellenv)"
elif [ -x /usr/local/bin/brew ]; then
  eval "$(/usr/local/bin/brew shellenv)"
fi
brew list git >/dev/null 2>&1 || brew install git
brew list node@22 >/dev/null 2>&1 || brew install node@22
brew list nginx >/dev/null 2>&1 || brew install nginx
export PATH="$(brew --prefix node@22)/bin:$PATH"
node_major="$(node -p 'Number(process.versions.node.split(".")[0])')"
[ "$node_major" -ge 22 ] || fail "Node 版本不足：$(node -v)，需要 22.x 或更高。"

say "2/9 获取或更新代码：$DEPLOY_DIR"
mkdir -p "$(dirname "$DEPLOY_DIR")"
if [ -d "$DEPLOY_DIR/.git" ]; then
  git -C "$DEPLOY_DIR" fetch --all --prune
  git -C "$DEPLOY_DIR" pull --ff-only || warn "git pull 失败，保留当前代码继续尝试。"
else
  git clone "$REPO_URL" "$DEPLOY_DIR"
fi
cd "$DEPLOY_DIR"

say "3/9 修正目标机部署路径"
export DEPLOY_DIR
find scripts -maxdepth 1 -type f -name "*.sh" -print0 \
  | xargs -0 perl -0pi -e 's#/Users/billchow/Documents/智慧冷冻站#$ENV{DEPLOY_DIR}#g'
chmod +x scripts/*.sh || true

say "4/9 安装前端和 BFF 依赖"
npm --prefix apps/chiller-bff ci
npm --prefix apps/chiller-shell-v1 ci
npm --prefix apps/chiller-admin-v1 ci

say "5/9 写入目标机本地环境文件"
cat > apps/chiller-shell-v1/.env.local <<EOF
VITE_BFF_BASE_URL=http://127.0.0.1:${BFF_PORT}
VITE_LEGACY_BASE_URL=${LEGACY_BASE_URL}
VITE_SITE_ID=${SITE_ID}
VITE_TREND_RANGE=24h
VITE_APP_MODE=${APP_MODE}
VITE_APP_MODE_LABEL=${APP_MODE_LABEL}
VITE_APP_READ_ONLY=${READ_ONLY_MODE}
VITE_UI_BADGE_STATE_URL=/ui-badge-state-v1.8.json
EOF
cat > apps/chiller-admin-v1/.env.local <<EOF
VITE_ADMIN_API_BASE_URL=http://127.0.0.1:${BFF_PORT}
VITE_LEGACY_BASE_URL=${LEGACY_BASE_URL}
VITE_ADMIN_APP_MODE=${APP_MODE}
VITE_ADMIN_APP_MODE_LABEL=${APP_MODE_LABEL}
VITE_ADMIN_READ_ONLY=${READ_ONLY_MODE}
VITE_ADMIN_USE_MOCKS=false
VITE_ADMIN_DEV_LOGIN=true
VITE_ADMIN_DEFAULT_SITE_ID=${SITE_ID}
VITE_ADMIN_BOOTSTRAP_USERNAMES=admin
EOF

say "6/9 初始化新项目配置库"
mkdir -p "$(dirname "$ADMIN_DB_FILE")"
ADMIN_DB_FILE="$ADMIN_DB_FILE" SITE_ID="$SITE_ID" \
  node apps/chiller-bff/scripts/ensure-energy-demo-subsystems.js \
  --profile=b25-cold-only \
  --site-id="$SITE_ID"
ADMIN_DB_FILE="$ADMIN_DB_FILE" SITE_ID="$SITE_ID" SITE_NAME="$SITE_NAME" \
  npm --prefix apps/chiller-bff run admin:ensure-runtime-config -- \
  --site-id="$SITE_ID" \
  --site-name="$SITE_NAME" \
  --min-condenser-inlet-temp-c="$MIN_CONDENSER_INLET_TEMP_C" \
  --tower-approach-dispatch-mode=shadow

say "7/9 准备 PLC Gateway"
if [ "$PLC_GATEWAY_MODE" = "mock" ]; then
  mkdir -p .local
  cat > .local/plc-gateway-mock.mjs <<'NODE'
import http from "node:http";
const port = Number(process.env.PLC_GATEWAY_PORT || 8877);
function json(res, status, body) {
  res.writeHead(status, {"content-type": "application/json"});
  res.end(JSON.stringify(body));
}
const server = http.createServer((req, res) => {
  if (req.method === "GET" && req.url === "/healthz") {
    json(res, 200, {ok: true, mode: "mock", plcWrite: false});
    return;
  }
  if (req.method === "POST" && (req.url === "/plc/dispatch/tower-approach" || req.url === "/plc/rollback/tower-approach")) {
    let raw = "";
    req.on("data", (chunk) => { raw += chunk; });
    req.on("end", () => {
      const payload = raw ? JSON.parse(raw) : {};
      const command = payload.command || {};
      json(res, 200, {
        ok: true,
        accepted: true,
        mode: "mock",
        plcWrite: false,
        writtenPoint: req.url.includes("rollback") ? "TCWS_SP_ROLLBACK" : "TCWS_SP",
        writtenValue: command.targetTcwsC ?? command.rollbackTarget?.targetTcwsC ?? null,
        readbackValue: command.targetTcwsC ?? command.rollbackTarget?.targetTcwsC ?? null,
        plcMode: "simulated-auto",
        interlockOk: true,
        alarmBlocked: false
      });
    });
    return;
  }
  json(res, 404, {ok: false, error: "not_found"});
});
server.listen(port, "127.0.0.1", () => console.log(`mock PLC gateway on ${port}`));
NODE
  launchctl remove "com.chiller.plc-gateway.mock.${PLC_GATEWAY_PORT}" >/dev/null 2>&1 || true
  : > "/tmp/chiller-plc-gateway-${PLC_GATEWAY_PORT}.log"
  launchctl submit \
    -l "com.chiller.plc-gateway.mock.${PLC_GATEWAY_PORT}" \
    -o "/tmp/chiller-plc-gateway-${PLC_GATEWAY_PORT}.log" \
    -e "/tmp/chiller-plc-gateway-${PLC_GATEWAY_PORT}.log" \
    -- /bin/zsh -lc "cd \"$DEPLOY_DIR\" && export PLC_GATEWAY_PORT=\"$PLC_GATEWAY_PORT\" PATH=\"$(brew --prefix node@22)/bin:\$PATH\" && exec node .local/plc-gateway-mock.mjs"
elif [ "$PLC_GATEWAY_MODE" = "real" ]; then
  curl -fsS "http://127.0.0.1:${PLC_GATEWAY_PORT}/healthz" >/tmp/chiller-plc-health.json \
    || fail "真实 PLC Gateway 未就绪：http://127.0.0.1:${PLC_GATEWAY_PORT}/healthz"
else
  warn "PLC_GATEWAY_MODE=$PLC_GATEWAY_MODE，跳过 Gateway 启动。"
fi

say "8/9 启动 BFF + 前台 + 管理端"
APP_MODE="$APP_MODE" BFF_PORT="$BFF_PORT" FRONTEND_PORT="$FRONTEND_PORT" \
  "$DEPLOY_DIR/scripts/stop_local_stack.sh" || true
APP_MODE="$APP_MODE" APP_MODE_LABEL="$APP_MODE_LABEL" READ_ONLY_MODE="$READ_ONLY_MODE" \
  BFF_PORT="$BFF_PORT" FRONTEND_PORT="$FRONTEND_PORT" SITE_ID="$SITE_ID" \
  LEGACY_BASE_URL="$LEGACY_BASE_URL" REALTIME_PARAMS_BASE_URL="$REALTIME_PARAMS_BASE_URL" \
  REALTIME_PARAMS_TIMEOUT_MS="$REALTIME_PARAMS_TIMEOUT_MS" \
  "$DEPLOY_DIR/scripts/start_local_stack.sh"

launchctl remove "com.chiller.admin.${APP_MODE}.${ADMIN_PORT}" >/dev/null 2>&1 || true
: > "/tmp/chiller-admin-${APP_MODE}-${ADMIN_PORT}.log"
launchctl submit \
  -l "com.chiller.admin.${APP_MODE}.${ADMIN_PORT}" \
  -o "/tmp/chiller-admin-${APP_MODE}-${ADMIN_PORT}.log" \
  -e "/tmp/chiller-admin-${APP_MODE}-${ADMIN_PORT}.log" \
  -- /bin/zsh -lc "cd \"$DEPLOY_DIR/apps/chiller-admin-v1\" && export PATH=\"$(brew --prefix node@22)/bin:\$PATH\" VITE_ADMIN_API_BASE_URL=\"http://127.0.0.1:${BFF_PORT}\" VITE_ADMIN_USE_MOCKS=false VITE_ADMIN_DEV_LOGIN=true VITE_ADMIN_DEFAULT_SITE_ID=\"$SITE_ID\" && exec npm run dev -- --host 127.0.0.1 --port \"$ADMIN_PORT\""

wait_http() {
  local name="$1"; local url="$2"; local i
  for i in $(seq 1 30); do
    if curl -fsS "$url" >/dev/null 2>&1; then
      printf '[OK] %s: %s\n' "$name" "$url"
      return 0
    fi
    sleep 1
  done
  return 1
}
wait_http "BFF" "http://127.0.0.1:${BFF_PORT}/healthz" || fail "BFF 未就绪"
wait_http "前台" "http://127.0.0.1:${FRONTEND_PORT}/" || fail "前台未就绪"
wait_http "管理端" "http://127.0.0.1:${ADMIN_PORT}/" || fail "管理端未就绪"
if [ "$PLC_GATEWAY_MODE" != "none" ]; then
  wait_http "PLC Gateway" "http://127.0.0.1:${PLC_GATEWAY_PORT}/healthz" || fail "PLC Gateway 未就绪"
fi

say "9/9 验收检查"
npm --prefix apps/chiller-bff run check:contract
set +e
SITE_ID="$SITE_ID" BFF_BASE_URL="http://127.0.0.1:${BFF_PORT}" FRONTEND_BASE_URL="http://127.0.0.1:${FRONTEND_PORT}" \
  "$DEPLOY_DIR/scripts/check_stack.sh" "$SITE_ID"
stack_rc=$?
set -e
if [ "$stack_rc" -ne 0 ]; then
  warn "check_stack 返回非零，通常是旧系统/现场上游未通。安装是否成功以 BFF/前台/管理端 health 为准。"
fi

if [ "$ENABLE_REAL_PLC_WRITE" = "1" ]; then
  if [ "$I_UNDERSTAND_REAL_PLC_WRITE" != "YES" ]; then
    fail "禁止打开真实写入：必须同时设置 ENABLE_REAL_PLC_WRITE=1 和 I_UNDERSTAND_REAL_PLC_WRITE=YES。"
  fi
  curl -fsS "http://127.0.0.1:${PLC_GATEWAY_PORT}/healthz" >/dev/null
  ADMIN_DB_FILE="$ADMIN_DB_FILE" SITE_ID="$SITE_ID" \
    npm --prefix apps/chiller-bff run admin:ensure-runtime-config -- \
    --site-id="$SITE_ID" \
    --tower-approach-dispatch-mode=enforced \
    --tower-approach-dispatch-approve-endpoint="http://127.0.0.1:${PLC_GATEWAY_PORT}/plc/dispatch/tower-approach" \
    --tower-approach-dispatch-rollback-endpoint="http://127.0.0.1:${PLC_GATEWAY_PORT}/plc/rollback/tower-approach" \
    --tower-approach-dispatch-timeout-ms=8000
  warn "已切 tower approach enforced 配置。仍需现场完成点位签字、PLC 联锁、回读和回退验收后才能真实运行。"
fi

cat <<EOF

=== 安装完成 ===
前台:      http://127.0.0.1:${FRONTEND_PORT}
管理端:    http://127.0.0.1:${ADMIN_PORT}
BFF:       http://127.0.0.1:${BFF_PORT}/healthz
PLC网关:   http://127.0.0.1:${PLC_GATEWAY_PORT}/healthz
项目:      ${SITE_ID} / ${SITE_NAME}
配置库:    ${ADMIN_DB_FILE}
默认边界:  READ_ONLY_MODE=${READ_ONLY_MODE}, PLC_GATEWAY_MODE=${PLC_GATEWAY_MODE}
日志:      /tmp/chiller-*.log

EOF
'''


CODEX_PROMPT = r'''你现在在一台全新的目标 Mac 上，不是在原开发电脑。请按以下要求完成“智慧冷冻站新项目本地平台一键安装”。

目标：
1. 在目标 Mac 部署 BFF、前台 Shell、管理端和 PLC Gateway mock。
2. 默认只启用 read-only + shadow，不允许真实 PLC 写入。
3. 安装完成后输出 URL、healthz 结果、关键日志路径和验收结论。

必须使用的变量：
REPO_URL=<替换成真实 git 仓库地址>
SITE_ID=new_chiller_site_001
SITE_NAME=新冷站项目
DEPLOY_DIR=$HOME/Documents/智慧冷冻站
PLC_GATEWAY_MODE=mock
ENABLE_REAL_PLC_WRITE=0

执行要求：
1. 先检查目标机系统、Node、npm、brew、端口占用。
2. 在 /tmp/install_chiller_target_mac.sh 写入本文档“完整一键安装脚本”。
3. 执行：
   REPO_URL="$REPO_URL" SITE_ID="$SITE_ID" SITE_NAME="$SITE_NAME" DEPLOY_DIR="$DEPLOY_DIR" bash /tmp/install_chiller_target_mac.sh
4. 若脚本失败，不要改安全边界；先读 /tmp/chiller-*.log 和终端报错，修复路径、端口、依赖或仓库问题。
5. 最终回报必须包含：
   - 前台 URL 是否 2xx
   - 管理端 URL 是否 2xx
   - BFF /healthz JSON 摘要
   - PLC Gateway /healthz JSON 摘要
   - npm --prefix apps/chiller-bff run check:contract 是否通过
   - 是否仍处于 READ_ONLY_MODE=1 和 shadow
   - 未完成项：真实 PLC 点表、写入白名单、回读、联锁、回退验收

安全边界：
- 不要把 mock PLC Gateway 当成真实 PLC。
- 不要把安装成功等同于真实闭环成功。
- 未经现场书面确认，不得设置 ENABLE_REAL_PLC_WRITE=1。
- pump-delta-t 不进入无人值守 enforced；冷却塔接近度也必须经过 PLC 本地保护、回读和回退验收。'''


def set_font(run, name="Calibri", east_asia="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def paragraph_border_bottom(paragraph, color="2E74B5", size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def set_para_spacing(paragraph, before=0, after=6, line=1.25):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_p(doc, text="", style=None, bold=False, color=INK, size=11, after=6):
    p = doc.add_paragraph(style=style)
    set_para_spacing(p, after=after)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, color=color, bold=bold)
    return p


def add_heading(doc, text, level=1):
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt({1: 18, 2: 14, 3: 10}.get(level, 8))
    p.paragraph_format.space_after = Pt({1: 10, 2: 7, 3: 5}.get(level, 4))
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}.get(level, 11), color=BLUE if level < 3 else DARK_BLUE, bold=True)
    return p


def add_callout(doc, title, body, fill=FILL_BLUE, tone=INK):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para_spacing(p, after=2)
    r = p.add_run(title)
    set_font(r, size=10.5, color=tone, bold=True)
    p2 = cell.add_paragraph()
    set_para_spacing(p2, after=0, line=1.2)
    r2 = p2.add_run(body)
    set_font(r2, size=10.5, color=INK)
    add_p(doc, "", after=3)


def add_kv_table(doc, rows, widths=(2160, 7200), header=None):
    extra = 1 if header else 0
    table = doc.add_table(rows=len(rows) + extra, cols=2)
    set_table_geometry(table, list(widths))
    start = 0
    if header:
        cell = table.cell(0, 0)
        cell.merge(table.cell(0, 1))
        set_cell_shading(cell, FILL_BLUE)
        p = cell.paragraphs[0]
        set_para_spacing(p, after=0)
        r = p.add_run(header)
        set_font(r, size=10.5, bold=True, color=DARK_BLUE)
        start = 1
    for ridx, (label, value) in enumerate(rows, start=start):
        c0, c1 = table.row_cells(ridx)
        set_cell_shading(c0, FILL_GRAY)
        p0 = c0.paragraphs[0]
        set_para_spacing(p0, after=0)
        r0 = p0.add_run(label)
        set_font(r0, size=9.5, color=DARK_BLUE, bold=True)
        p1 = c1.paragraphs[0]
        set_para_spacing(p1, after=0, line=1.18)
        r1 = p1.add_run(value)
        set_font(r1, size=9.5, color=INK)
    add_p(doc, "", after=3)
    return table


def add_matrix_table(doc, headers, rows, widths):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, FILL_BLUE)
        p = cell.paragraphs[0]
        set_para_spacing(p, after=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_font(r, size=9, color=DARK_BLUE, bold=True)
    for r_i, row in enumerate(rows, start=1):
        for c_i, value in enumerate(row):
            cell = table.cell(r_i, c_i)
            p = cell.paragraphs[0]
            set_para_spacing(p, after=0, line=1.15)
            if c_i == 0:
                set_cell_shading(cell, "FAFBFC")
            r = p.add_run(str(value))
            set_font(r, size=8.8, color=INK, bold=(c_i == 0))
            if c_i in (1, 3):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_p(doc, "", after=3)
    return table


def add_code_block(doc, code, title=None):
    if title:
        p = add_p(doc, title, bold=True, color=DARK_BLUE, size=10.5, after=4)
        p.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, FILL_CODE)
    cell.text = ""
    for idx, line in enumerate(code.rstrip().splitlines()):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line if line else " ")
        set_font(r, name="Courier New", east_asia="SimSun", size=7.2, color=RGBColor(22, 27, 34))
    add_p(doc, "", after=3)


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.375 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)


def number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)


def setup_document(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    for style_name in ["Normal", "Body Text"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.font.color.rgb = INK
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.25

    for idx, size in [(1, 16), (2, 13), (3, 12)]:
        style = styles[f"Heading {idx}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE if idx < 3 else DARK_BLUE

    header = section.header.paragraphs[0]
    set_para_spacing(header, after=0)
    r = header.add_run("智慧冷冻站新项目本地部署 | 目标 Mac Codex 一键安装")
    set_font(r, size=9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_para_spacing(footer, after=0)
    rf = footer.add_run("部署手册 | 只读 / shadow 默认边界")
    set_font(rf, size=9, color=MUTED)


def build_doc():
    doc = Document()
    setup_document(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("新冷站项目目标 Mac Codex 一键安装部署手册")
    set_font(r, size=24, color=RGBColor(0, 0, 0), bold=True)
    p2 = doc.add_paragraph()
    set_para_spacing(p2, after=10)
    r2 = p2.add_run("适用于另一台全新 Mac 本地平台；可联网安装；目标是让目标主机 Codex 直接照文档执行。")
    set_font(r2, size=12.5, color=MUTED)
    meta = doc.add_paragraph()
    set_para_spacing(meta, after=12)
    rm = meta.add_run(f"版本日期：{date.today().isoformat()}    默认边界：READ_ONLY_MODE=1 + shadow    系统范围：冷站优先")
    set_font(rm, size=10.5, color=MUTED, bold=True)
    paragraph_border_bottom(meta)

    add_callout(
        doc,
        "关键结论",
        "这份文档可以直接交给目标 Mac 上的 Codex：先粘贴第 2 章提示词，再让目标 Codex 写入并执行第 3 章脚本。脚本默认只部署本地平台和 mock PLC Gateway，不会真实写 PLC；真实闭环必须额外打开双重确认并完成现场联锁、回读、回退验收。",
        fill=FILL_WARN,
        tone=RISK,
    )

    add_heading(doc, "1. 使用方式", 1)
    number(doc, "在目标 Mac 打开 Codex，新建一个线程。")
    number(doc, "把第 2 章“直接喂给目标主机 Codex 的提示词”全文粘贴进去。")
    number(doc, "把真实仓库地址替换到 REPO_URL；没有真实 PLC 网关时保持 PLC_GATEWAY_MODE=mock。")
    number(doc, "让目标 Codex 执行脚本并返回 healthz、URL、合同检查和日志摘要。")
    number(doc, "只有现场点表、写入白名单、PLC 联锁、回读和回退验收完成后，才允许讨论 ENABLE_REAL_PLC_WRITE=1。")

    add_kv_table(
        doc,
        [
            ("目标系统", "另一台 Mac；可联网；用 Homebrew 安装 git、node@22、nginx。"),
            ("部署目录", "$HOME/Documents/智慧冷冻站，脚本会把仓库中的旧绝对路径替换为目标机路径。"),
            ("运行端口", "前台 3001，管理端 3002，BFF 8787，PLC Gateway 8877。"),
            ("默认边界", "READ_ONLY_MODE=1，towerApproachDispatchMode=shadow，mock Gateway 不做真实 PLC 写入。"),
            ("适用系统", "冷站优先：冷机、冷冻泵、冷却泵、冷却塔、COP、AI 优化建议、执行单治理。"),
        ],
        header="目标机部署基线",
    )

    add_heading(doc, "2. 直接喂给目标主机 Codex 的提示词", 1)
    add_p(doc, "将下面整段粘贴给目标 Mac 上的新 Codex 线程。目标 Codex 不需要知道当前电脑路径，只需要替换 REPO_URL。", color=MUTED)
    add_code_block(doc, CODEX_PROMPT, "Codex 提示词")

    add_heading(doc, "3. 完整一键安装脚本", 1)
    add_p(
        doc,
        "目标 Codex 应把本节脚本写入 /tmp/install_chiller_target_mac.sh 后执行。脚本包含依赖安装、克隆仓库、路径修正、npm ci、项目初始化、mock PLC Gateway、BFF/前台/管理端启动和基础验收。",
    )
    add_code_block(doc, INSTALL_SCRIPT, "install_chiller_target_mac.sh")

    add_heading(doc, "4. 部署架构", 1)
    add_matrix_table(
        doc,
        ["层级", "端口", "职责", "验收信号"],
        [
            ("前台 Shell", "3001", "运行端页面、总览、趋势、优化演示、执行单审阅。", "HTTP 2xx；/optimize-demo?siteId=... 可打开。"),
            ("BFF", "8787", "聚合站点数据、AI Advisor、执行单、管理端 API、合同检查。", "/healthz 返回 ok，check:contract 通过。"),
            ("管理端", "3002", "站点配置、子系统能力、runtime-config、点位映射复核。", "VITE_ADMIN_USE_MOCKS=false 下可进入登录/配置页。"),
            ("PLC Gateway", "8877", "HTTP 到 PLC 协议的隔离层；mock 模式仅模拟回读。", "/healthz 返回 mode=mock 或 real。"),
            ("PLC / SCADA", "现场网络", "真实点位、联锁、手自动、回读、回退。", "现场签字后才允许真实写入。"),
        ],
        [1800, 900, 4050, 2610],
    )

    add_heading(doc, "5. 安装变量说明", 1)
    add_matrix_table(
        doc,
        ["变量", "默认值", "是否必改", "说明"],
        [
            ("REPO_URL", "空", "是", "真实 git 仓库地址；脚本为空会停止。"),
            ("DEPLOY_DIR", "$HOME/Documents/智慧冷冻站", "建议确认", "目标机部署目录。"),
            ("SITE_ID", "new_chiller_site_001", "建议改", "目标项目唯一站点 ID。"),
            ("SITE_NAME", "新冷站项目", "建议改", "管理端展示名称。"),
            ("PLC_GATEWAY_MODE", "mock", "否", "mock / real / none；默认不会真实写 PLC。"),
            ("READ_ONLY_MODE", "1", "否", "BFF 写保护总闸。"),
            ("ENABLE_REAL_PLC_WRITE", "0", "现场前禁止改", "真实写 PLC 触发变量，必须另配确认短语。"),
            ("I_UNDERSTAND_REAL_PLC_WRITE", "NO", "现场前禁止改", "只有值为 YES 且 ENABLE_REAL_PLC_WRITE=1 才进入 enforced 配置。"),
        ],
        [2100, 2100, 1450, 3710],
    )

    add_heading(doc, "6. 目标 Codex 执行后的验收标准", 1)
    add_matrix_table(
        doc,
        ["检查项", "命令", "通过标准"],
        [
            ("BFF 健康", "curl http://127.0.0.1:8787/healthz", "返回 JSON；appMode、readOnlyMode、legacyBaseUrl 可读。"),
            ("前台入口", "curl -I http://127.0.0.1:3001/", "HTTP 2xx；浏览器可打开。"),
            ("管理端入口", "curl -I http://127.0.0.1:3002/", "HTTP 2xx；VITE_ADMIN_USE_MOCKS=false。"),
            ("PLC Gateway", "curl http://127.0.0.1:8877/healthz", "mock 模式返回 plcWrite=false；real 模式必须有真实联锁摘要。"),
            ("合同检查", "npm --prefix apps/chiller-bff run check:contract", "通过；OpenAPI 示例与 schema 一致。"),
            ("本地栈检查", "SITE_ID=$SITE_ID scripts/check_stack.sh $SITE_ID", "BFF/前台通过；旧系统上游不通可作为 WARN/现场资源缺口。"),
        ],
        [1900, 3300, 4160],
    )

    doc.add_page_break()
    add_heading(doc, "7. 真实 PLC 闭环门禁", 1)
    add_callout(
        doc,
        "禁止误判",
        "安装成功、页面能打开、mock Gateway accepted、BFF ok=true，都不等于真实 PLC 闭环成功。真实闭环的最终真值是：PLC 写入回读、联锁状态、告警闭锁、回退记录和现场签字。",
        fill=FILL_WARN,
        tone=RISK,
    )
    add_matrix_table(
        doc,
        ["门禁", "冷却塔接近度", "泵温差修正"],
        [
            ("第一版策略", "可做受控 enforced 候选，但默认先 shadow。", "只做 assisted 小步修正，不做无人值守 enforced。"),
            ("写点要求", "AI 专用目标点或明确批准的 TCWS_SP；不得写手自动模式点。", "AI_ChwpFreqTrim_Hz / AI_CwpFreqTrim_Hz，必须是修正量而非 PID 覆盖。"),
            ("PLC 保护", "上下限、单步 <=0.5℃、死区、最小保持、告警闭锁、回退。", "频率上下限、斜率、最小流量、末端安全、回退后闭锁。"),
            ("回读验收", "目标写入值与回读值一致；异常时自动回退或冻结。", "修正量生效、TTL 到期归零、回退为 0Hz 或稳定值。"),
            ("现场资料", "点位表、权限等级、写风险清单、PLC 逻辑截图/导出。", "泵频反馈、末端阀位/压差/室温、最小流量保护确认。"),
        ],
        [1800, 3780, 3780],
    )

    add_heading(doc, "8. 故障处理", 1)
    add_matrix_table(
        doc,
        ["现象", "优先检查", "处理"],
        [
            ("REPO_URL 未设置", "脚本开头变量", "补真实 git 地址后重跑。"),
            ("Node 版本不足", "node -v", "brew install node@22，并把 node@22 加入 PATH。"),
            ("3001/3002/8787 占用", "lsof -iTCP:<port> -sTCP:LISTEN", "停旧进程或改 FRONTEND_PORT / ADMIN_PORT / BFF_PORT。"),
            ("BFF 起不来", "/tmp/chiller-bff-*.log", "看 npm、路径、ADMIN_DB_FILE、端口和语法错误。"),
            ("前台空白", "/tmp/chiller-shell-*.log 与 VITE_BFF_BASE_URL", "确认 BFF health，确认 .env.local 已写目标端口。"),
            ("check_stack 失败", "失败项是 legacy、bff 还是 frontend", "legacy 上游不通是现场资源问题；BFF/前台失败才是安装问题。"),
            ("mock Gateway accepted", "/tmp/chiller-plc-gateway-*.log", "只能说明模拟链路通，不代表真实 PLC 写入。"),
        ],
        [2000, 3000, 4360],
    )

    add_heading(doc, "9. 目标 Codex 最终回报模板", 1)
    add_code_block(
        doc,
        r'''安装结果：
- 部署目录：
- Git commit：
- SITE_ID / SITE_NAME：
- 前台 URL：
- 管理端 URL：
- BFF /healthz 摘要：
- PLC Gateway /healthz 摘要：
- check:contract：
- check_stack：
- 当前安全边界：READ_ONLY_MODE=1/0，towerApproachDispatchMode=shadow/enforced，PLC_GATEWAY_MODE=mock/real
- 真实 PLC 写入：未启用 / 已启用（若已启用，必须附现场签字和回读证据）
- 日志路径：
- 仍缺资源：
  1. 真实 PLC 点位表与写入白名单
  2. 手自动/联锁/告警闭锁确认
  3. 回读一致性与回退测试
  4. 30-60 分钟 shadow 对比数据
''',
        "交付回报格式",
    )

    add_heading(doc, "10. 现场交付注意事项", 1)
    bullet(doc, "不要让目标 Codex 为了“安装成功”擅自安装 Docker、改系统网络策略或绕过现场安全流程。")
    bullet(doc, "点位表只能证明地址和权限，不足以证明 assisted/enforced readiness；必须结合实时趋势、回读、联锁和回退。")
    bullet(doc, "冷却塔手自动点属于高风险模式点，默认禁写；优先写 AI 专用目标点或经过批准的目标设定点。")
    bullet(doc, "泵温差优化第一版保留人工确认和 PLC 本地保护，不能让 AI 覆盖原 PID 或直接启停泵。")
    bullet(doc, "验收口径要分清：安装验收、本地联调验收、shadow 验收、真实 PLC 闭环验收。")

    doc.add_page_break()
    add_heading(doc, "附录 A. 常用维护命令", 1)
    add_code_block(
        doc,
        r'''# 查看服务端口
lsof -iTCP:3001 -sTCP:LISTEN
lsof -iTCP:3002 -sTCP:LISTEN
lsof -iTCP:8787 -sTCP:LISTEN
lsof -iTCP:8877 -sTCP:LISTEN

# 停止本地 BFF + 前台
APP_MODE=site-local BFF_PORT=8787 FRONTEND_PORT=3001 ./scripts/stop_local_stack.sh

# 重新启动本地 BFF + 前台
APP_MODE=site-local READ_ONLY_MODE=1 SITE_ID=$SITE_ID BFF_PORT=8787 FRONTEND_PORT=3001 ./scripts/start_local_stack.sh

# BFF 合同检查
npm --prefix apps/chiller-bff run check:contract

# 查看日志
tail -n 80 /tmp/chiller-bff-*.log
tail -n 80 /tmp/chiller-shell-*.log
tail -n 80 /tmp/chiller-admin-*.log
tail -n 80 /tmp/chiller-plc-gateway-*.log
''',
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT.resolve())
